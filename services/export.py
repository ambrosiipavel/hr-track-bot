"""Выгрузка базы в Excel: два листа — кандидаты и история статусов."""

from datetime import datetime
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy.ext.asyncio import AsyncSession

from database.db import get_all_candidates, get_all_history
from database.models import ROLE_LABELS, STATUS_LABELS

HEADER_FILL = PatternFill("solid", fgColor="16171B")
HEADER_FONT = Font(color="FFFFFF", bold=True)


def _write_header(sheet, titles: list[str]) -> None:
    sheet.append(titles)
    for cell in sheet[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(vertical="center")
    sheet.freeze_panes = "A2"


def _autosize(sheet, max_width: int = 45) -> None:
    for column in sheet.columns:
        longest = max(len(str(cell.value or "")) for cell in column)
        letter = get_column_letter(column[0].column)
        sheet.column_dimensions[letter].width = min(longest + 3, max_width)


async def build_workbook(session: AsyncSession) -> tuple[bytes | None, str]:
    """Возвращает готовый файл в байтах и имя для отправки."""
    candidates = await get_all_candidates(session)
    if not candidates:
        return None, ""

    workbook = Workbook()

    # ---- лист 1: кандидаты ----
    sheet = workbook.active
    sheet.title = "Кандидаты"
    _write_header(sheet, [
        "№", "Фамилия", "Имя", "Телефон", "Должность", "Статус",
        "Координатор", "Добавил", "Создан", "Закрыт", "Дней в работе",
        "Причина отказа",
    ])

    for c in candidates:
        sheet.append([
            c.id, c.last_name, c.first_name, c.phone, c.position,
            STATUS_LABELS[c.status],
            c.coordinator.full_name if c.coordinator else "",
            c.created_by.full_name if c.created_by else "",
            c.created_at.strftime("%d.%m.%Y %H:%M"),
            c.closed_at.strftime("%d.%m.%Y %H:%M") if c.closed_at else "",
            c.days_open,
            c.reject_reason or "",
        ])
    _autosize(sheet)

    # ---- лист 2: история ----
    history_sheet = workbook.create_sheet("История статусов")
    _write_header(history_sheet, [
        "Кандидат №", "Кандидат", "Кто менял", "Роль",
        "Было", "Стало", "Комментарий", "Когда",
    ])

    for record in await get_all_history(session):
        history_sheet.append([
            record.candidate_id,
            record.candidate.full_name if record.candidate else "",
            record.user.full_name if record.user else "",
            ROLE_LABELS[record.user.role] if record.user else "",
            STATUS_LABELS[record.old_status] if record.old_status else "—",
            STATUS_LABELS[record.new_status],
            record.comment or "",
            record.created_at.strftime("%d.%m.%Y %H:%M"),
        ])
    _autosize(history_sheet)

    buffer = BytesIO()
    workbook.save(buffer)
    filename = f"candidates_{datetime.now():%Y-%m-%d}.xlsx"
    return buffer.getvalue(), filename
